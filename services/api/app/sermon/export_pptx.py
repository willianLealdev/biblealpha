"""Slides da ministração (PowerPoint .pptx) para projetar: título, uma tela por seção com os pontos
principais e uma tela para cada texto bíblico. Fundo escuro e títulos dourados (identidade do app).
Notas pessoais não entram nos slides."""

import io
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

from app.models import Sermon

BG = RGBColor(0x0A, 0x08, 0x06)
GOLD = RGBColor(0xE8, 0xB3, 0x3A)
TEXT = RGBColor(0xF3, 0xEC, 0xDF)
MUTED = RGBColor(0xA8, 0x9F, 0x8F)
W, H = Emu(12192000), Emu(6858000)  # 16:9
_MARK = re.compile(r"\*\*|==|\*")


def _clean(text: str) -> str:
    return _MARK.sub("", text).strip()


def _slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # em branco
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG
    return slide


def _box(
    slide,
    text: str,
    top: float,
    height: float,
    size: int,
    color: RGBColor,
    bold=False,
    italic=False,
    center=False,
):
    box = slide.shapes.add_textbox(
        Emu(int(W * 0.07)), Emu(int(H * top)), Emu(int(W * 0.86)), Emu(int(H * height))
    )
    frame = box.text_frame
    frame.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.alignment = 2 if center else 1  # 2 = centro, 1 = esquerda
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
        run.font.name = "Georgia"
    return box


def slides_for(sermon: Sermon) -> list[dict]:
    """Estrutura dos slides (também usada pelo modo Apresentar do app)."""
    out = [{"kind": "title", "title": sermon.title, "subtitle": sermon.main_reference or ""}]
    current: dict | None = None
    for block in sermon.content.get("blocks", []):
        kind = block.get("type")
        if kind == "section":
            current = {"kind": "section", "title": block.get("title") or "", "points": []}
            out.append(current)
        elif kind == "text":
            lines = [
                _clean(re.sub(r"^([-*•]|\d+[.)])\s+", "", ln)) for ln in (block.get("text") or "").split("\n")
            ]
            lines = [ln for ln in lines if ln]
            if current is None:
                current = {"kind": "section", "title": "", "points": []}
                out.append(current)
            current["points"] += lines
        elif kind == "scripture":
            verses = " ".join(f"{v['verse']} {v['text']}" for v in block.get("verses", []))
            out.append(
                {
                    "kind": "scripture",
                    "title": f"{block.get('reference')} ({block.get('translation')})",
                    "text": verses,
                }
            )
    # Seção com muitos pontos vira mais de uma tela
    final: list[dict] = []
    for s in out:
        if s["kind"] == "section" and len(s["points"]) > 5:
            for i in range(0, len(s["points"]), 5):
                final.append({**s, "points": s["points"][i : i + 5]})
        elif s["kind"] != "section" or s["title"] or s["points"]:
            final.append(s)
    return final


def build_pptx(sermon: Sermon) -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    for s in slides_for(sermon):
        slide = _slide(prs)
        if s["kind"] == "title":
            _box(slide, s["title"], 0.32, 0.25, 44, GOLD, bold=True, center=True)
            if s["subtitle"]:
                _box(slide, s["subtitle"], 0.58, 0.1, 24, MUTED, center=True)
        elif s["kind"] == "section":
            _box(slide, s["title"], 0.08, 0.15, 34, GOLD, bold=True)
            if s["points"]:
                _box(slide, "\n".join(f"• {p}" for p in s["points"]), 0.27, 0.65, 24, TEXT)
        else:
            _box(slide, s["title"], 0.08, 0.12, 30, GOLD, bold=True)
            size = 26 if len(s["text"]) < 400 else 20 if len(s["text"]) < 800 else 16
            _box(slide, s["text"], 0.24, 0.7, size, TEXT, italic=True)
        _box(slide, "BibleAlpha", 0.93, 0.05, 10, MUTED, center=True)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
