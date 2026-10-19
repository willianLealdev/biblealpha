"""Exportação da ministração para Word (.docx)."""

import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from app.models import Sermon

GOLD = RGBColor(0x8F, 0x64, 0x18)
GRAY = RGBColor(0x6B, 0x63, 0x55)
_INLINE = re.compile(r"(\*\*[^*]+\*\*|==[^=]+==|\*[^*\n]+\*)")


def _runs(paragraph, text: str) -> None:
    pos = 0
    for m in _INLINE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos : m.start()])
        token = m.group(0)
        if token.startswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("=="):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            run.font.color.rgb = GOLD
        else:
            paragraph.add_run(token[1:-1]).italic = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def build_docx(sermon: Sermon, labels: dict[str, str]) -> bytes:
    doc = Document()
    doc.styles["Normal"].font.name = "Georgia"
    doc.styles["Normal"].font.size = Pt(12)
    doc.add_heading(sermon.title, level=0)

    meta = [
        (labels["main"], sermon.main_reference),
        (labels["theme"], sermon.theme),
        (labels["objective"], sermon.objective),
        (labels["audience"], sermon.audience),
        (labels["event"], sermon.event),
        (labels["date"], sermon.scheduled_for.strftime("%d/%m/%Y") if sermon.scheduled_for else None),
        (labels["duration"], f"{sermon.duration_minutes} min" if sermon.duration_minutes else None),
    ]
    for label, value in meta:
        if value:
            p = doc.add_paragraph()
            p.add_run(f"{label}: ").bold = True
            p.add_run(str(value))

    for block in sermon.content.get("blocks", []):
        kind = block.get("type")
        if kind == "section":
            doc.add_heading(block.get("title") or "", level=2)
        elif kind == "text":
            for line in (block.get("text") or "").split("\n"):
                stripped = line.strip()
                if not stripped:
                    continue
                if re.match(r"^[-*•]\s+", stripped):
                    _runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^[-*•]\s+", "", stripped))
                elif re.match(r"^\d+[.)]\s+", stripped):
                    _runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+[.)]\s+", "", stripped))
                elif stripped.startswith("#"):
                    doc.add_heading(stripped.lstrip("#").strip(), level=3)
                else:
                    _runs(doc.add_paragraph(), stripped)
        elif kind == "scripture":
            head = doc.add_paragraph()
            run = head.add_run(f"{block.get('reference')} ({block.get('translation')})")
            run.bold = True
            run.font.color.rgb = GOLD
            body = doc.add_paragraph()
            body.paragraph_format.left_indent = Pt(18)
            for v in block.get("verses", []):
                num = body.add_run(f"{v['verse']} ")
                num.font.size = Pt(8)
                num.bold = True
                body.add_run(f"{v['text']} ").italic = True
        elif kind == "note":
            p = doc.add_paragraph()
            run = p.add_run(f"{labels['note']}: {block.get('text', '')}")
            run.italic = True
            run.font.color.rgb = GRAY

    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("BibleAlpha")
    run.font.size = Pt(8)
    run.font.color.rgb = GRAY
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
